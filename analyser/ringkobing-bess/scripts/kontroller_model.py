"""Uafhængig Python-replika af Excel-modellen. Bruges til at kontrollere formlerne og til enkeltfølsomhederne i Resumé.

Brug: python3 kontroller_model.py <genberegnet model.xlsx>
Filen skal være genberegnet (fx åbnet og gemt i Excel eller konverteret med LibreOffice), så parameterværdierne kan læses.
"""
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], data_only=True)
par = {}
for row in wb["Forudsætninger"].iter_rows(min_row=5, values_only=True):
    if row[0] and row[3] is not None and not isinstance(row[3], str): par[row[0]] = row[3:6]
bt = {}
ws = wb["Backtest"]
for row in ws.iter_rows(values_only=True):
    if row and row[0] in (1, 2, 4) and isinstance(row[2], (int, float)) and isinstance(row[3], (int, float)): bt[row[0]] = (row[2], row[3])

def irr(cf):
    f = lambda r: sum(c / (1 + r) ** i for i, c in enumerate(cf))
    lo, hi = -0.9, 1.0
    if f(lo) * f(hi) > 0: return float("nan")
    for _ in range(200):
        m = (lo + hi) / 2
        if f(lo) * f(m) <= 0: hi = m
        else: lo = m
    return m

def run(p):
    g, mwh = bt[int(p["dur"])]
    capex_pl = p["mw"] * 1000 * p["ckw"] + p["mw"] * p["dur"] * 1000 * p["ckwh"]; capex = capex_pl + p["dev"] + p["conn"]
    tc = (p["t_en"] + p["t_dso"] + p["t_afg"]) * 10 / p["fx"]; td = (p["t_enp"] + p["t_dsop"]) * 10 / p["fx"]
    pre, post = [-capex], [-capex]; saldo = capex; loss = 0
    for t in range(1, 21):
        on = 1 if t <= p["life"] else 0
        idx = (1 + p["infl"]) ** (p["cod"] + t - 1 - 2026); cf = (1 - p["fade"]) ** (t - 1)
        arb = p["mw"] * g * p["lvl"] * p["cap"] * (1 + p["idu"]) * (1 + p["arbtr"]) ** (t - 1) * cf * p["avail"] * idx * on
        anc = p["mw"] * max(p["anc1"] * (1 + p["anctr"]) ** (t - 1), p["ancfl"]) * p["avail"] * idx * on
        rev = arb + anc
        grid = -p["mw"] * mwh * cf * p["avail"] * (tc / p["rte"] + td) * idx * on
        e = rev + grid - (p["om"] + p["ins"]) * capex_pl * idx * on - p["fee"] * rev - p["rfv"] * p["mw"] * idx * on - p["fix"] * idx * on
        dep = (saldo if t == p["life"] else saldo * p["saldo"]) * on; saldo -= dep
        ti = e - dep; lu = min(loss, max(ti, 0)); loss = loss - lu + max(-ti, 0)
        pre.append(e); post.append(e - max(ti - lu, 0) * p["tax"])
    return capex, irr(pre), irr(post)

names = ["Pessimistisk", "Base", "Optimistisk"]
res = {r[0]: r[1:] for r in wb["Resumé"].iter_rows(min_row=5, max_row=16, max_col=4, values_only=True)}
for i, n in enumerate(names):
    c, a, b = run({k: v[i] for k, v in par.items()})
    print(f"{n}: replika IRR før/efter skat {a:.4%} / {b:.4%}; Excel {res['IRR før skat (ugearet)'][i]} / {res['IRR efter skat (ugearet)'][i]}")
B = {k: v[1] for k, v in par.items()}
tests = {"Systemydelser 70k": dict(anc1=70000), "Ingen betaling til RFV": dict(rfv=0), "Intraday 20 %": dict(idu=0.2),
         "Capex 250 EUR/kWh": dict(ckw=160, ckwh=170, dev=150000), "Levetid 20 år": dict(life=20), "4 timer": dict(dur=4),
         "Fuld elafgift": dict(t_afg=60.9)}
for k, v in tests.items():
    c, a, b = run(dict(B, **v)); print(f"Base + {k}: IRR efter skat {b:.1%}")
lo, hi = 0.2, 1.5
for _ in range(60):
    m = (lo + hi) / 2
    c, a, b = run(dict(B, ckw=B["ckw"] * m, ckwh=B["ckwh"] * m, dev=B["dev"] * m))
    lo, hi = (m, hi) if b > 0.08 else (lo, m)
print(f"Break-even investering for 8 % efter skat: {c / (B['mw'] * B['dur'] * 1000):.0f} EUR/kWh")
lo, hi = 0.5, 3
for _ in range(60):
    m = (lo + hi) / 2
    c, a, b = run(dict(B, lvl=B["lvl"] * m, anc1=B["anc1"] * m, ancfl=B["ancfl"] * m))
    lo, hi = (lo, m) if b > 0.08 else (m, hi)
print(f"Break-even markedsindtægt for 8 % efter skat: faktor {m:.2f}")
